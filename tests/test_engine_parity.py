"""Backend parity for the new scene features: SVG, PPTX and PNG (layout-engine STY-04, roadmap S2)."""
import re

import pytest

from spinagram import Diagram, Theme, render

THEME = Theme.load("default")


@pytest.fixture
def icon(tmp_path):
    p = tmp_path / "i.svg"
    p.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><rect width="10" height="10" '
                 'fill="#1F6FEB"/></svg>')
    return str(p)


def _body_font_theme(tmp_path):
    """A theme whose fonts have files without symbol glyphs, so fallback runs appear."""
    import string

    import yaml

    from test_engine_glyphs import make_font  # pytest puts tests/ on sys.path (rootdir conftest-less layout)
    f = make_font(tmp_path / "plain.ttf", string.printable.strip() + " ", "Plain")
    (tmp_path / "theme.yaml").write_text("extends: default\n" + yaml.safe_dump(
        {"fonts": {r: {"family": "Plain", "files": {"400": str(f), "700": str(f)}} for r in ("heading", "body")}}))
    return Theme.load(tmp_path)


def feature(name, d, icon):
    if name == "dash":
        d.zone("Net", x=200, y=200, w=800, h=400, dash="dashed", stroke="#202020", stroke_width=3, fill="none")
    elif name == "icon-only":
        d.box(None, id="chip", x=200, y=200, w=120, h=60, icon=icon, align="center")
    elif name == "classed-label":
        d.css = ".hdr { color: #AA0000 }"
        d.label("NETWORK", id="hdr", x=200, y=300, class_="hdr")
    elif name == "absolute-edge":
        d.edge((200, 300), (900, 300), "to a point")
    elif name == "zone-target":
        d.box("Pool", id="pool", x=200, y=300, w=200, h=100)
        z = d.zone("Net", id="net", x=700, y=200, w=500, h=300)
        d.edge("pool", z)
    elif name == "fallback-glyph":
        d.box("Status", id="n", x=200, y=200, w=500, h=160, text="ok ⚠ done")


SVG_CHECKS = {
    "dash": lambda svg: 'stroke-dasharray="8 6"' in svg,
    "icon-only": lambda svg: "<image" in svg and re.search(r'<rect[^>]*width="120\.0"[^>]*fill="[^"]*"(?![^>]*stroke=)', svg),
    "classed-label": lambda svg: re.search(r'<text[^>]*fill="#AA0000"[^>]*>.*NETWORK', svg),
    "absolute-edge": lambda svg: re.search(r"200\.0[ ,]300\.0", svg) and "to a point" in svg,
    "zone-target": lambda svg: re.search(r"700\.0[ ,]350\.0|700[ ,]350", svg),
    "fallback-glyph": lambda svg: re.search(r"<tspan font-family=\"&apos;Noto Sans Symbols 2&apos;[^\"]*\">⚠", svg),
}


def _pptx_check(name, prs):
    from pptx.oxml.ns import qn
    slide = prs.slides[0]
    xml = slide._element.xml
    shapes = {s.name: s for s in slide.shapes}
    if name == "dash":
        return "<a:custDash>" in xml
    if name == "icon-only":
        grp = shapes["chip (group)"]
        inner = {s.name: s for s in grp.shapes}
        ln = inner["chip"]._element.spPr.find(qn("a:ln"))
        return "chip icon" in inner and (ln is None or ln.find(qn("a:noFill")) is not None)
    if name == "classed-label":
        return any(s.has_text_frame and s.text_frame.text == "NETWORK" and
                   str(s.text_frame.paragraphs[0].runs[0].font.color.rgb) == "AA0000" for s in slide.shapes)
    if name == "absolute-edge":
        return "<p:cxnSp>" in xml and any(s.has_text_frame and s.text_frame.text == "to a point" for s in slide.shapes)
    if name == "zone-target":
        return f'<a:endCxn id="{shapes["z0"].shape_id}"' in xml
    if name == "fallback-glyph":
        return re.search(r'<a:sym typeface="Noto Sans Symbols 2"', xml) is not None


FEATURES = list(SVG_CHECKS)


def _render(tmp_path, name, icon, fmt):
    theme = _body_font_theme(tmp_path) if name == "fallback-glyph" else THEME
    d = Diagram("P", "Parity")
    feature(name, d, icon)
    return render([d], tmp_path / fmt, theme=theme, formats=[fmt])


@pytest.mark.parametrize("name", FEATURES)
def test_svg_parity(tmp_path, icon, name):
    svg = next(p for p in _render(tmp_path, name, icon, "svg") if p.suffix == ".svg").read_text()
    assert SVG_CHECKS[name](svg)


@pytest.mark.parametrize("name", FEATURES)
def test_pptx_parity(tmp_path, icon, name):
    pptx = pytest.importorskip("pptx")
    out = next(p for p in _render(tmp_path, name, icon, "pptx") if p.suffix == ".pptx")
    assert _pptx_check(name, pptx.Presentation(str(out)))


def _png_or_skip(tmp_path, name, icon):
    pytest.importorskip("playwright")
    try:
        return next(p for p in _render(tmp_path, name, icon, "png") if p.suffix == ".png")
    except Exception as exc:  # noqa: BLE001 - no browser installed
        if "Executable doesn't exist" in str(exc) or "playwright install" in str(exc):
            pytest.skip("chromium not installed for playwright")
        raise


@pytest.mark.parametrize("name", FEATURES)
def test_png_renders_every_feature(tmp_path, icon, name):
    from PIL import Image
    with Image.open(_png_or_skip(tmp_path, name, icon)) as im:
        assert im.size == (1920 * 2, 1080 * 2)


@pytest.mark.parametrize("dash,dashed", [("dashed", True), ("solid", False)])
def test_png_draws_the_dash_pattern(tmp_path, icon, dash, dashed):
    from PIL import Image
    d = Diagram("P", "Parity")
    d.zone("Net", x=200, y=200, w=800, h=400, dash=dash, stroke="#202020", stroke_width=3, fill="none")
    pytest.importorskip("playwright")
    try:
        png = next(p for p in render([d], tmp_path, theme=THEME, formats=["png"]) if p.suffix == ".png")
    except Exception as exc:  # noqa: BLE001
        if "playwright install" in str(exc) or "Executable doesn't exist" in str(exc):
            pytest.skip("chromium not installed for playwright")
        raise
    with Image.open(png) as im:
        row = [im.getpixel((x, 400))[:3] for x in range(500, 1900)]   # the top outline at 2x scale
    dark = [sum(p) < 200 for p in row]
    transitions = sum(1 for a, b in zip(dark, dark[1:]) if a != b)
    assert (transitions > 20) == dashed and any(dark)
