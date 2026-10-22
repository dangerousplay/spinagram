"""Edge endpoints: node anchors, zone targets, bus fans (layout-engine EDG-07..12)."""
import pytest

from spinagram import Diagram, build_scene, Theme
from spinagram.scene import Connector

THEME = Theme.load("default")

# box "b" at (400, 400, 200, 100); the other end is an absolute point
ANCHOR_SCENARIOS = [
    # id, b endpoint, other point, expected point on b, expected binding (side, frac) or None
    ("auto-right", "b", (900, 450), (600, 450), ("right", 0.5)),
    ("auto-left", "b", (100, 450), (400, 450), ("left", 0.5)),
    ("auto-top", "b", (500, 100), (500, 400), ("top", 0.5)),
    ("auto-bottom", "b", (500, 900), (500, 500), ("bottom", 0.5)),
    ("side", "b.top", (900, 900), (500, 400), ("top", 0.5)),
    ("side-frac", "b.left@0.25", (0, 0), (400, 425), ("left", 0.25)),
    ("center-unbound", "b.center", (0, 0), (500, 450), None),
    ("absolute", (123, 456), (0, 0), (123, 456), None),
]


def _edge_scene(src, dst):
    d = Diagram("EP", "Endpoints")
    d.box("B", id="b", x=400, y=400, w=200, h=100)
    d.edge(src, dst)
    return [i for i in build_scene(d, THEME).items if isinstance(i, Connector)][0]


@pytest.mark.parametrize("ref,other,point,binding", [s[1:] for s in ANCHOR_SCENARIOS],
                         ids=[s[0] for s in ANCHOR_SCENARIOS])
def test_node_anchor_resolution(ref, other, point, binding):
    c = _edge_scene(ref, other)
    assert c.points[0] == pytest.approx(point)
    got = (c.src.side, c.src.frac) if c.src else None
    assert got == binding


def test_unknown_node_raises():
    with pytest.raises(KeyError, match="unknown node 'nope'"):
        _edge_scene("nope", (0, 0))


# zone endpoints (EDG-07, EDG-08, EDG-12) ----------------------------------------------------------------------
# zone "net" at (400, 400, 400, 300): centre (600, 550); points stay 24 px from corners
ZONE_SCENARIOS = [
    # id, source point, zone ref, expected end point, expected side
    ("left", (100, 550), "zone:net", (400, 550), "left"),
    ("right", (1100, 450), "zone:net", (800, 450), "right"),
    ("above", (600, 100), "zone:net", (600, 400), "top"),
    ("below", (600, 1000), "zone:net", (600, 700), "bottom"),
    ("top-left-corner-clamped", (100, 100), "zone:net", (424, 400), "top"),
    ("bottom-right-corner-clamped", (1100, 1000), "zone:net", (776, 700), "bottom"),
    ("left-clamped-low", (100, 690), "zone:net", (400, 676), "left"),
    ("inside-inner-face", (450, 600), "zone:net", (400, 600), "left"),
    ("explicit-side-frac", (0, 0), "zone:net.top@0.25", (500, 400), "top"),
]


def _zone_edge(src, dst, **kw):
    d = Diagram("EP", "Endpoints")
    z = d.zone("Net", id="net", x=400, y=400, w=400, h=300)
    d.edge(src, z if dst == "ZONE" else dst, **kw)
    return [i for i in build_scene(d, THEME).items if isinstance(i, Connector)][0]


@pytest.mark.parametrize("src,ref,point,side", [s[1:] for s in ZONE_SCENARIOS], ids=[s[0] for s in ZONE_SCENARIOS])
def test_zone_endpoint_side_and_point(src, ref, point, side):
    c = _zone_edge(src, ref)
    assert c.points[-1] == pytest.approx(point)
    assert (c.dst.shape_id, c.dst.side) == ("z0", side)


def test_zone_object_and_node_source():
    d = Diagram("EP", "Endpoints")
    z = d.zone("Net", id="net", x=400, y=400, w=400, h=300)
    pool = d.box("Pool", id="pool", x=1000, y=500, w=200, h=100)
    d.edge(pool, z)
    c = [i for i in build_scene(d, THEME).items if isinstance(i, Connector)][0]
    assert c.points == [pytest.approx((1000, 550)), pytest.approx((800, 550))]
    assert (c.src.shape_id, c.src.side, c.dst.shape_id, c.dst.side) == ("pool", "left", "z0", "right")


def test_zone_endpoint_from_yaml():
    from spinagram.yaml_io import from_dict
    d = from_dict({"id": "EP", "title": "E", "zones": [{"id": "net", "x": 400, "y": 400, "w": 400, "h": 300}],
                   "edges": ["zone:net.left -> zone:net.right"]})
    c = [i for i in build_scene(d, THEME).items if isinstance(i, Connector)][0]
    assert c.points == [pytest.approx((400, 550)), pytest.approx((800, 550))]


def test_unknown_zone_raises():
    with pytest.raises(KeyError, match="unknown zone 'nope'"):
        _zone_edge((0, 0), "zone:nope")


@pytest.mark.parametrize("text,expected", [
    ("pool -> zone:net", ("pool", "zone:net", None)),
    ("pool -> zone:net.top@0.3: HTTPS", ("pool", "zone:net.top@0.3", "HTTPS")),
    ("zone:a.right -> b: x", ("zone:a.right", "b", "x")),
], ids=["zone-target", "zone-target-label", "zone-source"])
def test_edge_shorthand_with_zone_refs(text, expected):
    from spinagram.model import parse_edge
    e = parse_edge(text)
    assert (e.src, e.dst, e.label) == expected


# through-box exclusion and PPTX binding (EDG-10, EDG-11) ------------------------------------------------------
def _through(edges):
    from spinagram.qa import check
    d = Diagram("TB", "Through")
    d.box("Pool", id="pool", x=100, y=400, w=200, h=100)
    z = d.zone("Net", id="net", x=600, y=300, w=700, h=300)
    d.box("Node A", id="a", x=650, y=400, w=200, h=100)
    d.box("Node B", id="b", x=1000, y=400, w=200, h=100)
    for src, dst in edges:
        d.edge(src, z if dst == "ZONE" else dst)
    return sorted((i.item, i.detail) for i in check(build_scene(d, THEME)) if i.kind == "edge-through-box")


THROUGH_SCENARIOS = [
    ("to-zone-border", [("pool", "ZONE")], []),
    ("to-zone-side", [("pool", "zone:net.left")], []),
    ("to-first-node", [("pool", "a")], []),
    ("through-sibling", [("pool", "b")], [("e0", "edge passes through 'a'")]),
    ("absolute-through", [((50, 450), (1300, 450))], [("e0", "edge passes through 'a'"),
                                                      ("e0", "edge passes through 'b'"),
                                                      ("e0", "edge passes through 'pool'")]),
]


@pytest.mark.parametrize("edges,expected", [s[1:] for s in THROUGH_SCENARIOS], ids=[s[0] for s in THROUGH_SCENARIOS])
def test_edge_through_box_excludes_zone_targets(edges, expected):
    assert _through(edges) == expected


@pytest.mark.parametrize("ref,point", [("ZONE", (600, 450)), ("zone:net.top@0.25", (775, 300))],
                         ids=["auto", "side-frac"])
def test_pptx_connector_bound_to_zone_at_exact_point(tmp_path, ref, point):
    pptx = pytest.importorskip("pptx")
    from pptx.oxml.ns import qn
    from spinagram import render
    from spinagram.backends.pptx import EMU
    d = Diagram("PB", "Bind")
    d.box("Pool", id="pool", x=100, y=400, w=200, h=100)
    z = d.zone("Net", id="net", x=600, y=300, w=700, h=300)
    d.edge("pool", z if ref == "ZONE" else ref)
    render([d], tmp_path, formats=["pptx"])
    shapes = {s.name: s for s in pptx.Presentation(str(tmp_path / "diagrams.pptx")).slides[0].shapes}
    cx = next(s for s in shapes.values() if s.shape_type is not None and "Connector" in type(s).__name__)
    end = cx._element.find(".//" + qn("a:endCxn"))
    assert end is not None and end.get("id") == str(shapes["z0"].shape_id)
    px = cx.end_x / EMU, cx.end_y / EMU
    assert px == pytest.approx(point, abs=0.5)


# bus fan (EDG-09) ----------------------------------------------------------------------------------------------
from spinagram.scene import Line  # noqa: E402

# zone (600, 300, 600, 400); children are 100x60 boxes laid along the facing side
BUS_SIDES = [
    # id, source point, side, bar axis offset (10 px inside the side)
    ("left", (100, 500), "left", 610),
    ("top", (900, 50), "top", 310),
    ("right", (1500, 500), "right", 1190),
    ("bottom", (900, 900), "bottom", 690),
]


def _bus_scene(src, n_children):
    d = Diagram("BUS", "Bus")
    z = d.zone("Net", id="net", x=600, y=300, w=600, h=400)
    kids = []
    for k in range(n_children):  # a diagonal so that both axes spread
        kids.append(d.box(f"N{k}", id=f"n{k}", x=700 + k * 150, y=380 + k * 100, w=100, h=60))
    d.edge(src, z, fan="bus")
    sc = build_scene(d, THEME)
    c = [i for i in sc.items if isinstance(i, Connector)][0]
    bars = [i for i in sc.items if isinstance(i, Line) and i.id == "e0-bus"]
    return c, bars, kids


@pytest.mark.parametrize("n", [0, 1, 3], ids=lambda n: f"{n}children")
@pytest.mark.parametrize("src,side,offset", [s[1:] for s in BUS_SIDES], ids=[s[0] for s in BUS_SIDES])
def test_bus_bar_inside_zone_spans_children(src, side, offset, n):
    c, bars, kids = _bus_scene(src, n)
    assert len(bars) == 1 and c.arrow_end and c.dst.side == side
    bar, end = bars[0], c.points[-1]
    horizontal = side in ("top", "bottom")
    if horizontal:
        assert bar.y1 == bar.y2 == offset
        lo, hi = bar.x1, bar.x2
        want = [k.x + k.w / 2 for k in kids] or [600 + 24, 1200 - 24]
        v = end[0]
    else:
        assert bar.x1 == bar.x2 == offset
        lo, hi = bar.y1, bar.y2
        want = [k.y + k.h / 2 for k in kids] or [300 + 24, 700 - 24]
        v = end[1]
    assert (lo, hi) == (min(want + [v]), max(want + [v]))
    assert end == pytest.approx({"left": (600, end[1]), "right": (1200, end[1]), "top": (end[0], 300),
                                 "bottom": (end[0], 700)}[side])  # arrowhead stays at the border


def test_no_bus_without_fan_or_zone_target():
    d = Diagram("BUS", "Bus")
    d.zone("Net", id="net", x=600, y=300, w=600, h=400)
    d.box("A", id="a", x=700, y=400, w=100, h=60)
    d.edge((100, 500), "zone:net")
    d.edge((100, 500), "a", fan="bus")
    assert not [i for i in build_scene(d, THEME).items if isinstance(i, Line) and i.id.endswith("-bus")]
