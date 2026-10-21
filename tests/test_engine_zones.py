"""Zones: model fields, outlines, CSS, headers (layout-engine ZON-01..13)."""
import pytest
import yaml

from spinagram import Diagram
from spinagram.yaml_io import dump, from_dict

ZONE_FIELD_SCENARIOS = [
    ("id", {"id": "net-a"}),
    ("dash-word", {"dash": "dotted"}),
    ("dash-list", {"dash": [4, 2]}),
    ("stroke", {"stroke": "primary", "stroke_width": 1.5}),
    ("fill-none", {"fill": "none"}),
    ("header", {"header": {"padding": 14, "strip": "primary"}}),
    ("icon-chips", {"icon": ["logos:vmware", "logos:redhat"]}),
    ("fit", {"fit": "children", "padding": 30, "equalize": "siblings"}),
    ("meaning", {"meaning": "network"}),
    ("class", {"cls": "az"}),
]


@pytest.mark.parametrize("kw", [k for _, k in ZONE_FIELD_SCENARIOS], ids=[i for i, _ in ZONE_FIELD_SCENARIOS])
def test_zone_fields_python_and_yaml_parity(kw):
    d = Diagram("Z", "Zones")
    z = d.zone("Net", x=10, y=20, w=300, h=200, **kw)
    for k, v in kw.items():
        assert getattr(z, k) == v
    ykw = {("class" if k == "cls" else k): v for k, v in kw.items()}
    y = from_dict({"id": "Z", "title": "Zones", "zones": [{"label": "Net", "x": 10, "y": 20, "w": 300, "h": 200, **ykw}]})
    assert y.zones[0] == z


def test_zone_ids_default_to_index_and_must_be_unique():
    d = Diagram("Z", "Zones")
    a = d.zone("A", x=0, y=0, w=10, h=10)
    b = d.zone("B", x=0, y=0, w=10, h=10, id="net")
    c = d.zone("C", x=0, y=0, w=10, h=10)
    assert (a.id, b.id, c.id) == ("z0", "net", "z2")
    with pytest.raises(ValueError, match="duplicate zone id 'net'"):
        d.zone("D", x=0, y=0, w=10, h=10, id="net")


@pytest.mark.parametrize("kw", [k for _, k in ZONE_FIELD_SCENARIOS], ids=[i for i, _ in ZONE_FIELD_SCENARIOS])
def test_zone_fields_survive_yaml_dump(kw):
    d = Diagram("Z", "Zones")
    z = d.zone("Net", x=10, y=20, w=300, h=200, **kw)
    again = from_dict(yaml.safe_load(yaml.safe_dump(dump(d))))
    assert again.zones[0] == z


# outlines (ZON-09, ZON-10) -----------------------------------------------------------------------------------
from spinagram import Theme, build_scene  # noqa: E402
from spinagram.backends.svg import to_svg  # noqa: E402
from spinagram.scene import Shape  # noqa: E402

THEME = Theme.load("default")
DASH_VALUES = [
    ("solid", "solid", None, None),
    ("dashed", "dashed", [8, 6], None),
    ("dotted", "dotted", [2, 6], "round"),
    ("custom", [4, 2, 1, 2], [4, 2, 1, 2], None),
]


def _zone_shape(theme=THEME, **kw):
    d = Diagram("Z", "Zones")
    d.zone("Net", x=10, y=20, w=300, h=200, **kw)
    sc = build_scene(d, theme)
    return next(i for i in sc.items if isinstance(i, Shape) and i.id == "z0"), to_svg(sc)


def _style_theme(tmp_path, style):
    import yaml
    (tmp_path / "theme.yaml").write_text("extends: default\n" + yaml.safe_dump(
        {"styles": {"zone": {"net": {"fill": "none", "label": "text-muted", "stroke": "primary", **style}}}}))
    return Theme.load(tmp_path)


@pytest.mark.parametrize("source", ["zone", "style"])
@pytest.mark.parametrize("value,dash,cap", [d[1:] for d in DASH_VALUES], ids=[d[0] for d in DASH_VALUES])
def test_zone_dash_values(tmp_path, source, value, dash, cap):
    if source == "zone":
        s, svg = _zone_shape(dash=value, stroke="primary")
    else:
        s, svg = _zone_shape(_style_theme(tmp_path, {"dash": value}), color="net")
    assert s.dash == dash and s.cap == cap and s.dashed == (dash is not None)
    if dash:
        assert f'stroke-dasharray="{" ".join(f"{v:g}" for v in dash)}"' in svg
    assert ('stroke-linecap="round"' in svg) == (cap == "round")


def test_zone_dash_beats_style_dash(tmp_path):
    s, _ = _zone_shape(_style_theme(tmp_path, {"dash": "dotted"}), color="net", dash="dashed")
    assert (s.dash, s.cap) == ([8, 6], None)


def test_legacy_dashed_style_keeps_default_pattern(tmp_path):
    s, svg = _zone_shape(_style_theme(tmp_path, {"dashed": True}), color="net")
    assert s.dashed and s.dash is None and 'stroke-dasharray="10 7"' in svg


STROKE_FILL = [
    ("stroke-token", {"stroke": "primary"}, {"stroke": THEME.color("primary"), "stroke_width": 2}),
    ("stroke-width", {"stroke": "#112233", "stroke_width": 1.5}, {"stroke": "#112233", "stroke_width": 1.5}),
    ("fill-token", {"fill": "primary-soft"}, {"fill": THEME.color("primary-soft")}),
    ("fill-none", {"fill": "none"}, {"fill": None, "opacity": 1.0}),
]


@pytest.mark.parametrize("kw,expected", [s[1:] for s in STROKE_FILL], ids=[s[0] for s in STROKE_FILL])
def test_zone_stroke_and_fill(kw, expected):
    s, svg = _zone_shape(**kw)
    assert {k: getattr(s, k) for k in expected} == expected
    if kw.get("fill") == "none":
        assert '<rect x="10.0" y="20.0" width="300.0" height="200.0" rx="26" fill="none"' in svg


def test_unknown_dash_word_raises():
    with pytest.raises(ValueError, match="zone dash 'wavy'"):
        _zone_shape(dash="wavy")


# PPTX parity (ZON-12) ------------------------------------------------------------------------------------------
@pytest.mark.parametrize("value,dash,cap", [d[1:] for d in DASH_VALUES], ids=[d[0] for d in DASH_VALUES])
def test_zone_dash_pptx_matches_svg(tmp_path, value, dash, cap):
    pptx = pytest.importorskip("pptx")
    from pptx.oxml.ns import qn

    from spinagram import render
    d = Diagram("Z", "Zones")
    d.zone("Net", x=10, y=20, w=300, h=200, dash=value, stroke="primary", stroke_width=2)
    render([d], tmp_path, formats=["pptx", "svg"])
    sh = next(s for s in pptx.Presentation(str(tmp_path / "diagrams.pptx")).slides[0].shapes if s.name == "z0")
    ln = sh._element.spPr.find(qn("a:ln"))
    cust = ln.find(qn("a:custDash"))
    if dash is None:
        assert cust is None
    else:
        pat = dash + (dash if len(dash) % 2 else [])
        got = [(int(ds.get("d")), int(ds.get("sp"))) for ds in cust.findall(qn("a:ds"))]
        assert got == [(round(a / 2 * 100000), round(b / 2 * 100000)) for a, b in zip(pat[0::2], pat[1::2])]
        svg = next(tmp_path.glob("*.svg")).read_text()
        assert f'stroke-dasharray="{" ".join(f"{v:g}" for v in dash)}"' in svg
    assert ln.get("cap") == ("rnd" if cap == "round" else None)


# CSS on zones (ZON-11, ZON-13, STY-01) ------------------------------------------------------------------------
CSS_SELECTORS = [
    ("type-id", "zone#net-a"),
    ("type-class", "zone.az"),
    ("type-style", "zone.soft"),
    ("bare-id", "#net-a"),
    ("bare-class", ".az"),
    ("type", "zone"),
]


@pytest.mark.parametrize("selector", [s for _, s in CSS_SELECTORS], ids=[i for i, _ in CSS_SELECTORS])
def test_css_reaches_zones_without_gid(selector):
    d = Diagram("Z", "Zones", css=f"{selector} {{ stroke: #112233; stroke-width: 3; stroke-dasharray: 2 6 }}")
    d.zone("Net", x=10, y=20, w=300, h=200, id="net-a", class_="az")
    d.zone("Other", x=400, y=20, w=300, h=200, id="other", color="warning")
    sc = build_scene(d, THEME)
    shapes = {i.id: i for i in sc.items if isinstance(i, Shape)}
    assert (shapes["z0"].stroke, shapes["z0"].stroke_width, shapes["z0"].dash) == ("#112233", 3, [2, 6])
    assert (shapes["z1"].stroke == "#112233") == (selector == "zone")


NONE_PAINT = [
    ("fill", "fill: none", "fill", None),
    ("stroke", "stroke: none", "stroke", None),
]


@pytest.mark.parametrize("decl,attr,expected", [s[1:] for s in NONE_PAINT], ids=[s[0] for s in NONE_PAINT])
def test_css_none_paint_on_zones(decl, attr, expected):
    d = Diagram("Z", "Zones", css=f"zone {{ {decl} }}")
    d.zone("Net", x=10, y=20, w=300, h=200, stroke="primary")
    s = next(i for i in build_scene(d, THEME).items if isinstance(i, Shape) and i.id == "z0")
    assert getattr(s, attr) is expected


def test_css_none_on_text_colour_warns_and_keeps_colour():
    from spinagram.qa import check
    d = Diagram("Z", "Zones", css="node { color: none }")
    d.box("A", id="a", x=0, y=0, w=200, h=100)
    sc = build_scene(d, THEME)
    a = next(i for i in sc.items if isinstance(i, Shape) and i.id == "a")
    assert a.label.blocks[0].color.startswith("#")
    assert any("color: none" in i.detail for i in check(sc))


# header (ZON-01, 02, 03, 07, 08) ----------------------------------------------------------------------------
from spinagram.scene import Text  # noqa: E402

@pytest.fixture
def icon_path(tmp_path):
    p = tmp_path / "icon.svg"
    p.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><rect width="10" height="10"/></svg>')
    return str(p)


def _header_scene(icon, label="Availability zone site-a", theme=THEME, **kw):
    d = Diagram("Z", "Zones")
    d.zone(label, x=100, y=200, w=600, h=300, icon=icon, **kw)
    d.zone("Second", x=800, y=200, w=400, h=300, icon=icon, dash="dotted", stroke="primary")
    sc = build_scene(d, theme)
    by = {i.id: i for i in sc.items if isinstance(i, (Shape, Text))}
    return by, sc


def _label_center(t):
    return t.top + t.block.leading / 2


HEADER_SCENARIOS = [
    pytest.param(dash, label, id=f"{dash}-{len(label)}ch")
    for dash in ("solid", "dashed", "dotted") for label in ("AZ", "Availability zone site-a · RHEL 9")]


@pytest.mark.parametrize("dash,label", HEADER_SCENARIOS)
def test_inset_header_geometry(icon_path, dash, label):
    by, _ = _header_scene(icon_path, label, dash=dash, stroke="primary")
    icon, lab = by["z0-icon"].icon, by["z0-label"]
    assert icon.x - 100 >= 14 and icon.y - 200 >= 14                  # ZON-01: padding inside the outline
    assert icon.size == by["z1-icon"].icon.size == 28                 # ZON-02: same size for every zone
    assert abs(_label_center(lab) - (icon.y + icon.size / 2)) <= 1    # ZON-03: label centred on the icon
    assert lab.x >= icon.x + icon.size                                # label after the icon


@pytest.mark.parametrize("source", ["zone", "style", "theme"])
@pytest.mark.parametrize("padding,size", [(14, 28), (20, 36)], ids=["default", "custom"])
def test_header_options_cascade(tmp_path, icon_path, source, padding, size):
    import yaml
    hdr = {"padding": padding, "icon_size": size}
    theme, kw = THEME, {}
    if source == "zone":
        kw = {"header": hdr}
    else:
        data = {"defaults": {"zone": {"header": hdr}}} if source == "theme" else \
            {"styles": {"zone": {"hdr": {"fill": "primary-soft", "label": "text-muted", "header": hdr}}}}
        (tmp_path / "theme.yaml").write_text("extends: default\n" + yaml.safe_dump(data))
        theme = Theme.load(tmp_path)
        kw = {"color": "hdr"} if source == "style" else {}
    by, _ = _header_scene(icon_path, theme=theme, **kw)
    icon = by["z0-icon"].icon
    assert (icon.x - 100, icon.y - 200, icon.size) == (padding, padding, size)


def test_flush_header_keeps_corner_icon(icon_path):  # ZON-07
    by, _ = _header_scene(icon_path, header={"inset": False})
    icon = by["z0-icon"].icon
    assert (icon.x, icon.y, icon.size) == (100, 200, 32)
    assert (by["z0-label"].x, by["z0-label"].y) == (100 + 32 + 10, 222)


def test_aws_styles_stay_flush():  # ZON-07: the AWS convention is kept by the theme
    for name in ("aws-cloud", "aws-account", "aws-region", "aws-vpc", "aws-subnet"):
        assert THEME.zone_style(name)["header"] == {"inset": False}


@pytest.mark.parametrize("upper,expected", [(True, "SITE-A"), (False, "site-a")], ids=["upper", "keep-case"])
def test_header_upper(icon_path, upper, expected):  # ZON-08
    by, _ = _header_scene(icon_path, "site-a", header={"upper": upper})
    assert by["z0-label"].block.lines == [expected]


def test_unresolved_icon_lays_out_label_alone():  # edge case
    by, _ = _header_scene("nopack:missing-icon")
    assert "z0-icon" not in by and (by["z0-label"].x, by["z0-label"].y) == (124, 236)


def test_zone_without_icon_keeps_label_place():
    by, _ = _header_scene(None)
    assert (by["z0-label"].x, by["z0-label"].y) == (124, 236)


# chips, strip, content_top (ZON-04, 05, 06) ------------------------------------------------------------------
@pytest.mark.parametrize("count", [1, 2, 3], ids=lambda n: f"{n}chips")
@pytest.mark.parametrize("strip", [None, "primary"], ids=["no-strip", "strip"])
def test_header_chips_and_strip(icon_path, count, strip):
    d = Diagram("Z", "Zones")
    d.zone("Site A", x=100, y=200, w=600, h=300, icon=[icon_path] * count, header={"strip": strip} if strip else None)
    sc = build_scene(d, THEME)
    by = {i.id: i for i in sc.items if isinstance(i, (Shape, Text))}
    chips = [by["z0-icon"]] + [by[f"z0-icon{k}"] for k in range(1, count)]
    xs = [c.icon.x for c in chips]
    assert xs == sorted(xs) and len(set(xs)) == count                           # ZON-04: left to right
    assert all(c.icon.y == chips[0].icon.y and c.icon.size == 28 for c in chips)
    assert by["z0-label"].x >= chips[-1].icon.x + 28                           # label after the last chip
    if strip:                                                                  # ZON-05: 5 px bar at the top
        bar = by["z0-strip"]
        assert (bar.y, bar.h, bar.fill) == (200, 5, THEME.color("primary")) and 100 <= bar.x < bar.x + bar.w <= 700
    else:
        assert "z0-strip" not in by


CONTENT_TOP = [
    ("icon", "ICON", 200 + 14 + 28 + 12),
    ("no-icon", None, 200 + 36 + 16 * 0.3 + 12),
]


@pytest.mark.parametrize("icon,expected", [s[1:] for s in CONTENT_TOP], ids=[s[0] for s in CONTENT_TOP])
def test_content_top(icon_path, icon, expected):  # ZON-06
    d = Diagram("Z", "Zones")
    z = d.zone("Site A", x=100, y=200, w=600, h=300, icon=icon_path if icon else None, id="site")
    assert z.content_top == pytest.approx(expected)
    assert build_scene(d, THEME).zone_content_tops["site"] == pytest.approx(expected)
