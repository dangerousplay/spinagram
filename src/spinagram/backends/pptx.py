"""PowerPoint with native, editable objects.

* nodes → autoshapes (rounded rectangle, diamond) with their text inside the shape;
  icons are grouped with their shape
* edges → connectors **bound** to their shapes when anchored at a side centre (moving a
  box drags the arrow); other routes become polylines; arrowheads via ``a:tailEnd``
* tables → native tables; diagram ``source`` → speaker notes
* SVG icons/images → pictures with a PNG fallback in ``a:blip`` and the ``asvg:svgBlip`` extension
  pointing to the SVG part (what PowerPoint writes on "Insert > Pictures"): vectors in PowerPoint
  2019+/365, the PNG elsewhere. Fallback: sibling ``.png`` → resvg render (4x, min 256 px, cached) →
  1x1 transparent. ``PptxBackend(svg="raster")`` or ``SPINAGRAM_PPTX_SVG=raster`` writes the PNG only.

One canvas pixel = 6350 EMU (0.5 pt), so a 1920x1080 canvas is a 13.333x7.5 in slide.
"""
from __future__ import annotations

import hashlib
import io
import os
import warnings
from pathlib import Path
from typing import Sequence

from .. import shapes
from ..scene import Connector, Image, Label, Line, Scene, Shape, TableGrid, Text, TextBlock

EMU = 6350
CXN_IDX = {"top": 0, "left": 1, "bottom": 2, "right": 3}
NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _require():
    try:
        import pptx  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("PowerPoint needs the 'pptx' extra: uv add 'spinagram[pptx]'") from exc


def _e(px: float) -> int:
    return int(round(px * EMU))


def _rgb(hex_: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_.lstrip("#")[:6].upper())


def _blend(hex_: str, bg: str, alpha: float) -> str:
    a = [int(hex_.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(bg.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
    return "#" + "".join(f"{round(x * alpha + y * (1 - alpha)):02X}" for x, y in zip(a, b))


SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
PNG_1X1 = bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
                        "0000000d49444154789c6360606060000000050001a5f645400000000049454e44ae426082")
_warned: set[str] = set()


def _warn_once(key: str, msg: str) -> None:
    if key not in _warned:
        _warned.add(key)
        warnings.warn(msg, stacklevel=3)


def _fallback(svg: Path, data: bytes, w: float, h: float):
    """PNG stand-in for an SVG: sibling .png → resvg render (4x, min 256 px, cached) → 1x1 transparent."""
    sib = svg.with_suffix(".png")
    if sib.exists():
        return str(sib)
    k = max(4.0, 256 / max(1.0, w, h))
    pw, ph = max(1, round(w * k)), max(1, round(h * k))
    from ..icons import cache_root

    cache = cache_root().parent / "vector" / f"pptx-{hashlib.sha256(data).hexdigest()}-{pw}x{ph}.png"
    if cache.exists():
        return str(cache)
    try:
        import resvg_py

        png = bytes(resvg_py.svg_to_bytes(svg_string=data.decode("utf-8"), width=pw, height=ph))
    except ImportError:
        _warn_once("resvg", "SVG pictures get a blank PNG fallback in PowerPoint: install the 'vector' "
                            "extra (resvg-py) for a real one")
        return io.BytesIO(PNG_1X1)
    except Exception as exc:  # noqa: BLE001 - a broken SVG must not break the deck
        _warn_once(f"resvg:{svg}", f"{svg.name}: PNG fallback render failed ({exc}); using a blank one")
        return io.BytesIO(PNG_1X1)
    try:
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(png)
    except OSError:
        return io.BytesIO(png)
    return str(cache)


def _segments(text: str, b: TextBlock) -> list[tuple[str, object]]:
    """``text`` cut where the block's fallback runs change font (one run when there is no fallback)."""
    if not b.runs:
        return [(text, b.font)]
    font_of = {ch: r.font for line in b.runs for r in line if r.font is not b.font for ch in r.text}
    out: list[list] = []
    for ch in text:
        f = font_of.get(ch, b.font) if not ch.isspace() else (out[-1][1] if out else b.font)
        if out and out[-1][1] is f:
            out[-1][0] += ch
        else:
            out.append([ch, f])
    return [(t, f) for t, f in out]


class _Slide:
    def __init__(self, prs, scene: Scene, svg: str = "vector", svg_parts: dict | None = None):
        self.scene = scene
        self.svg = svg
        self.svg_parts = {} if svg_parts is None else svg_parts  # sha256 → SVG Part (one per deck)
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.shapes: dict[str, object] = {}
        bg = self.slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(scene.background)
        notes = [scene.notes] if scene.notes else []
        if scene.sources:
            notes.append("Fontes dos números:")
            for r in scene.sources:
                where = " · ".join(str(r[k]) for k in ("system", "query", "formula", "computed_at") if r.get(k))
                notes.append(f"• {r['display']} — {r['label']} — {where}")
        if notes:
            self.slide.notes_slide.notes_text_frame.text = "\n".join(notes)

    # styling helpers ----------------------------------------------------------
    @staticmethod
    def _line(shape, color: str | None, width: float, dashed: bool = False, dash: list[float] | None = None,
              cap: str | None = None) -> None:
        from lxml import etree
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.oxml.ns import qn

        if not color or not width:
            shape.line.fill.background()
            return
        shape.line.color.rgb = _rgb(color)
        shape.line.width = _e(width)
        if dash:  # exact pattern: custDash lengths are in 1/1000 % of the line width
            ln = shape.line._get_or_add_ln()
            for old in ln.findall(qn("a:prstDash")) + ln.findall(qn("a:custDash")):
                ln.remove(old)
            cust = etree.SubElement(ln, qn("a:custDash"))
            pat = list(dash) + (list(dash) if len(dash) % 2 else [])
            for d, sp in zip(pat[0::2], pat[1::2]):
                ds = etree.SubElement(cust, qn("a:ds"))
                ds.set("d", str(round(d / width * 100000)))
                ds.set("sp", str(round(sp / width * 100000)))
            # schema order: fill, prstDash/custDash, join, headEnd, tailEnd
            for tag in ("a:round", "a:bevel", "a:miter", "a:headEnd", "a:tailEnd"):
                for el in ln.findall(qn(tag)):
                    ln.remove(el)
                    ln.append(el)
        elif dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if cap == "round":  # dotted outlines: round dots
            shape.line._get_or_add_ln().set("cap", "rnd")

    @staticmethod
    def _shadow(shape, sh: tuple | None) -> None:
        """CSS box-shadow → DrawingML outer shadow."""
        if not sh:
            return
        import math

        from lxml import etree
        from pptx.oxml.ns import qn

        dx, dy, blur, color, alpha = sh
        sp_pr = shape._element.spPr
        eff = etree.SubElement(sp_pr, qn("a:effectLst"))
        o = etree.SubElement(eff, qn("a:outerShdw"))
        o.set("blurRad", str(_e(blur)))
        o.set("dist", str(_e(math.hypot(dx, dy))))
        o.set("dir", str(round((math.degrees(math.atan2(dy, dx)) % 360) * 60000)))
        o.set("algn", "ctr")
        o.set("rotWithShape", "0")
        clr = etree.SubElement(o, qn("a:srgbClr"))
        clr.set("val", color.strip("#").upper())
        etree.SubElement(clr, qn("a:alpha")).set("val", str(round(alpha * 100000)))

    @staticmethod
    def _fill_alpha(shape, alpha: float) -> None:
        """CSS fill-opacity → alpha on the solid fill colour (the line keeps full opacity)."""
        if alpha >= 1:
            return
        from lxml import etree
        from pptx.oxml.ns import qn

        clr = shape._element.spPr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
        if clr is not None:
            etree.SubElement(clr, qn("a:alpha")).set("val", str(round(alpha * 100000)))

    _MARKER = {"arrow": "triangle", "open-arrow": "arrow", "diamond": "diamond", "open-diamond": "diamond",
               "circle": "oval", "bar": None}

    @classmethod
    def _arrows(cls, shape, start: bool, end: bool, kinds: tuple[str, str] = ("arrow", "arrow")) -> None:
        from pptx.oxml.ns import qn
        from lxml import etree

        ln = shape.line._get_or_add_ln()
        for tag, on, kind, prop in (("a:headEnd", start, kinds[0], "marker-start"),
                                    ("a:tailEnd", end, kinds[1], "marker-end")):
            if on:
                typ = cls._MARKER.get(kind, "triangle")
                if kind in ("open-diamond", "bar"):
                    _warn_once(f"pptx-marker:{kind}", f"pptx: {prop} '{kind}' has no PowerPoint arrowhead; "
                               f"drawn as {typ or 'no arrowhead'}")
                if typ is None:
                    continue
                el = etree.SubElement(ln, qn(tag))
                el.set("type", typ)
                el.set("w", "med")
                el.set("len", "med")

    @staticmethod
    def _no_shadow(shape) -> None:
        """Drop the theme style reference (it adds shadows); fill/line/text are explicit."""
        el = shape._element
        for style in el.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}style"):
            el.remove(style)

    @staticmethod
    def _runs(tf, blocks: list[TextBlock], align: str, gap: float, prewrapped: bool) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        first = True
        for bi, b in enumerate(blocks):
            paras = b.lines if prewrapped else b.paragraphs
            for pi, text in enumerate(paras):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
                p.line_spacing = round(b.leading / b.font.size, 2)
                if bi > 0 and pi == 0:
                    p.space_before = Pt(gap * 0.5)
                for seg, font in _segments(text, b):
                    r = p.add_run()
                    r.text = seg
                    f = r.font
                    f.name = font.family
                    f.size = Pt(b.font.size * 0.5)
                    f.bold = b.font.weight >= 600
                    f.italic = b.font.italic
                    f.color.rgb = _rgb(b.color)
                    if font is not b.font:  # glyph fallback: the symbol font for this run (a:latin + a:sym)
                        from pptx.oxml.ns import qn

                        rpr = r._r.get_or_add_rPr()
                        sym = rpr.find(qn("a:sym"))
                        if sym is None:
                            sym = rpr.makeelement(qn("a:sym"), {})
                            latin = rpr.find(qn("a:latin"))
                            latin.addnext(sym) if latin is not None else rpr.append(sym)
                        sym.set("typeface", font.family)

    def _text_frame(self, shape, s: Shape, label: Label, prewrapped: bool) -> None:
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

        tf = shape.text_frame
        tf.word_wrap = not prewrapped
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bx, by, bw, bh = label.box
        if s.kind == "diamond" or s.kind == "pill":
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        else:
            tf.margin_left = _e(bx - s.x)
            tf.margin_right = _e(max(0, s.x + s.w - (bx + bw)))
            tf.margin_top = _e(max(0, by - s.y))
            tf.margin_bottom = _e(max(0, s.y + s.h - (by + bh)))
        self._runs(tf, label.blocks, label.align, label.gap, prewrapped)

    def picture(self, shapes, path: Path, x: float, y: float, w: float, h: float, name: str):
        """Add a picture; an SVG gets a PNG fallback plus the ``asvg:svgBlip`` extension (vector)."""
        if path.suffix.lower() != ".svg":
            pic = shapes.add_picture(str(path), _e(x), _e(y), _e(w), _e(h))
            pic.name = name
            return pic
        if not path.exists():
            _warn_once(f"missing:{path}", f"{name}: picture {path} not found; skipped")
            return None
        data = path.read_bytes()
        pic = shapes.add_picture(_fallback(path, data, w, h), _e(x), _e(y), _e(w), _e(h))
        pic.name = name
        if self.svg == "raster":
            return pic
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.opc.package import Part
        from pptx.oxml.ns import qn

        sha = hashlib.sha256(data).hexdigest()
        part = self.svg_parts.get(sha)
        slide_part = self.slide.part
        if part is None:
            pkg = slide_part.package
            part = Part(pkg.next_partname("/ppt/media/svg%d.svg"), "image/svg+xml", pkg, data)
            self.svg_parts[sha] = part
        rid = slide_part.relate_to(part, RT.IMAGE)
        blip = pic._element.blipFill.find(qn("a:blip"))
        ext_lst = blip.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"), uri=SVG_EXT)
        etree.SubElement(ext, f"{{{ASVG_NS}}}svgBlip", nsmap={"asvg": ASVG_NS}).set(qn("r:embed"), rid)
        return pic

    # items --------------------------------------------------------------------
    def shape(self, s: Shape, container=None) -> object:
        from pptx.enum.shapes import MSO_SHAPE

        target = container if container is not None else self.slide.shapes
        if s.icon is not None and container is None:
            group = self.slide.shapes.add_group_shape()
            group.name = f"{s.id} (group)"
            sh = self.shape(s, group.shapes)
            ic = s.icon
            self.picture(group.shapes, ic.path, ic.x, ic.y, ic.width, ic.size, f"{s.id} icon")
            return sh
        if s.parts:
            return self.symbol(s, target)
        kind = {"diamond": MSO_SHAPE.DIAMOND}.get(s.kind)
        if kind is None:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if s.radius else MSO_SHAPE.RECTANGLE
        sh = target.add_shape(kind, _e(s.x), _e(s.y), _e(s.w), _e(s.h))
        sh.name = s.id
        if s.radius and s.kind != "diamond":
            sh.adjustments[0] = min(0.5, s.radius / max(1.0, min(s.w, s.h)))
        if s.fill:
            sh.fill.solid()
            fill = _blend(s.fill, self.scene.background, s.opacity) if s.opacity < 1 else s.fill
            sh.fill.fore_color.rgb = _rgb(fill)
            self._fill_alpha(sh, s.fill_opacity)
        else:
            sh.fill.background()
        self._line(sh, s.stroke, s.stroke_width, s.dashed, s.dash, s.cap)
        self._no_shadow(sh)
        self._shadow(sh, s.shadow)
        if s.label:
            self._text_frame(sh, s, s.label, prewrapped=s.kind in ("diamond", "pill"))
        self.shapes[s.id] = sh
        return sh

    def _paint(self, sh, s: Shape, paint: str) -> None:
        if paint == "stroke":
            sh.fill.background()
        else:
            color = s.stroke if paint == "solid" else s.fill
            if color:
                sh.fill.solid()
                fill = _blend(color, self.scene.background, s.opacity) if s.opacity < 1 else color
                sh.fill.fore_color.rgb = _rgb(fill)
                if paint != "solid":
                    self._fill_alpha(sh, s.fill_opacity)
            else:
                sh.fill.background()
        self._line(sh, s.stroke, s.stroke_width, s.dashed, s.dash)
        self._no_shadow(sh)
        if paint != "stroke":
            self._shadow(sh, s.shadow)

    def symbol(self, s: Shape, target) -> object:
        """Behavior symbol: preset geometry, a group of natives for multi-part symbols, else a freeform."""
        from pptx.enum.shapes import MSO_SHAPE

        plan = shapes.pptx_plan(s.kind)
        if plan == "group":  # final: ring + dot
            group = target.add_group_shape()
            group.name = f"{s.id} (group)"
            for part, scale in zip(s.parts, (0.92, 0.6)):
                r = min(s.w, s.h) / 2 * scale
                cx, cy = s.x + s.w / 2, s.y + s.h / 2
                sh = group.shapes.add_shape(MSO_SHAPE.OVAL, _e(cx - r), _e(cy - r), _e(2 * r), _e(2 * r))
                sh.name = f"{s.id} {part.paint}"
                self._paint(sh, s, part.paint)
            self.shapes[s.id] = group
            return group
        if plan == "freeform":
            pts = []
            for sub in shapes.parse(s.parts[0].d):
                pts.append(sub.start)
                for cmd, ps in sub.segments:
                    if cmd == "L":
                        pts += ps
                    else:  # flatten cubic segments
                        (x0, y0), ((x1, y1), (x2, y2), (x3, y3)) = pts[-1], ps
                        for k in range(1, 9):
                            t = k / 8
                            pts.append(((1 - t) ** 3 * x0 + 3 * (1 - t) ** 2 * t * x1 + 3 * (1 - t) * t * t * x2
                                        + t ** 3 * x3, (1 - t) ** 3 * y0 + 3 * (1 - t) ** 2 * t * y1
                                        + 3 * (1 - t) * t * t * y2 + t ** 3 * y3))
            fb = target.build_freeform(_e(pts[0][0]), _e(pts[0][1]), scale=1.0)
            fb.add_line_segments([(_e(x), _e(y)) for x, y in pts[1:]], close=True)
            sh = fb.convert_to_shape()
        else:
            sh = target.add_shape(getattr(MSO_SHAPE, plan), _e(s.x), _e(s.y), _e(s.w), _e(s.h))
            if plan == "ROUNDED_RECTANGLE":
                sh.adjustments[0] = min(0.5, {"state": 18, "action": 22}.get(s.kind, 18) / max(1.0, min(s.w, s.h)))
        sh.name = s.id
        self._paint(sh, s, s.parts[0].paint)
        if s.label:
            self._text_frame(sh, s, s.label, prewrapped=False)
        self.shapes[s.id] = sh
        return sh

    def connector(self, c: Connector) -> None:
        from pptx.enum.shapes import MSO_CONNECTOR

        def bindable(b) -> bool:  # zone borders bind at any fraction; nodes at side midpoints (connection sites)
            return (b is not None and (b.frac == 0.5 or b.target == "zone") and b.side in CXN_IDX
                    and b.shape_id in self.shapes)

        bound = c.route == "straight" and any(bindable(b) for b in (c.src, c.dst))
        if bound or c.route == "straight":
            (x1, y1), (x2, y2) = c.points[0], c.points[-1]
            cx = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _e(x1), _e(y1), _e(x2), _e(y2))
            for b, begin in ((c.src, True), (c.dst, False)):
                if bindable(b):
                    target = self.shapes[b.shape_id]
                    (cx.begin_connect if begin else cx.end_connect)(target, CXN_IDX[b.side])
                    if b.target == "zone":  # binding snaps to the side midpoint; keep the exact border point
                        px, py = (x1, y1) if begin else (x2, y2)
                        if begin:
                            cx.begin_x, cx.begin_y = _e(px), _e(py)
                        else:
                            cx.end_x, cx.end_y = _e(px), _e(py)
            shape = cx
        else:
            pts = c.points
            if c.route == "curve":
                (x0, y0), (qx, qy), (x2, y2) = pts
                n = 24
                pts = [((1 - t) ** 2 * x0 + 2 * (1 - t) * t * qx + t * t * x2,
                        (1 - t) ** 2 * y0 + 2 * (1 - t) * t * qy + t * t * y2) for t in (i / n for i in range(n + 1))]
            fb = self.slide.shapes.build_freeform(_e(pts[0][0]), _e(pts[0][1]), scale=1.0)
            fb.add_line_segments([(_e(x), _e(y)) for x, y in pts[1:]], close=False)
            shape = fb.convert_to_shape()
            shape.fill.background()
        shape.name = c.id
        self._no_shadow(shape)
        self._line(shape, c.color, c.width, c.dashed, c.dash)
        if c.opacity < 1:  # line transparency: alpha on the line colour
            from lxml import etree
            from pptx.oxml.ns import qn

            clr = shape.line._get_or_add_ln().find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
            if clr is not None:
                etree.SubElement(clr, qn("a:alpha")).set("val", str(round(c.opacity * 100000)))
        self._arrows(shape, c.arrow_start, c.arrow_end, (c.marker_start, c.marker_end))
        if c.label:
            self.shape(c.label)

    def text(self, t: Text) -> None:
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

        w = (t.width or 200) + 8
        x = {"middle": t.x - w / 2, "end": t.x - w + 4}.get(t.anchor, t.x - 2)
        h = t.block.height
        tb = self.slide.shapes.add_textbox(_e(x), _e(t.top), _e(w), _e(h))
        tb.name = t.id
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        self._runs(tf, [t.block], {"middle": "center"}.get(t.anchor, "left"), 0, prewrapped=True)
        if t.anchor == "end":
            from pptx.enum.text import PP_ALIGN

            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.RIGHT

    def line(self, ln: Line) -> None:
        from pptx.enum.shapes import MSO_CONNECTOR

        cx = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _e(ln.x1), _e(ln.y1), _e(ln.x2), _e(ln.y2))
        cx.name = ln.id
        self._no_shadow(cx)
        self._line(cx, ln.color, ln.width)

    def image(self, im: Image) -> None:
        self.picture(self.slide.shapes, im.path, im.x, im.y, im.w, im.h, im.id)

    def table(self, t: TableGrid) -> None:
        from lxml import etree
        from pptx.enum.text import MSO_ANCHOR
        from pptx.oxml.ns import qn

        rows, cols = len(t.cells), len(t.col_widths)
        gf = self.slide.shapes.add_table(rows, cols, _e(t.x), _e(t.y), _e(t.w), _e(t.h))
        gf.name = t.id
        tbl = gf.table
        tbl.first_row = False
        tbl.horz_banding = False
        style = gf._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = NO_STYLE_NO_GRID
        for i, w in enumerate(t.col_widths):
            tbl.columns[i].width = _e(w)
        y = t.y
        for r, (rh, row) in enumerate(zip(t.row_heights, t.cells)):
            tbl.rows[r].height = _e(rh)
            x = t.x
            for c, (cw, cell) in enumerate(zip(t.col_widths, row)):
                pc = tbl.cell(r, c)
                if cell.fill:
                    pc.fill.solid()
                    pc.fill.fore_color.rgb = _rgb(cell.fill)
                pc.margin_left = _e(cell.pad_left)
                pc.margin_right = _e(10)
                pc.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = pc.text_frame
                tf.word_wrap = True
                self._runs(tf, [cell.block], "left", 0, prewrapped=False)
                if r > 0:
                    tcPr = pc._tc.get_or_add_tcPr()
                    lnB = etree.Element(qn("a:lnB"), w=str(_e(1)))
                    sf = etree.SubElement(lnB, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=t.border.lstrip("#"))
                    tcPr.insert(0, lnB)
                if cell.icon:
                    s = cell.icon.size
                    self.picture(self.slide.shapes, cell.icon.path, x + 14, y + (rh - s) / 2, s, s,
                                 f"{t.id}[{r},{c}] icon")
                x += cw
            y += rh

    def build(self) -> None:
        for it in self.scene.items:
            if isinstance(it, Shape):
                self.shape(it)
            elif isinstance(it, Connector):
                self.connector(it)
            elif isinstance(it, Text):
                self.text(it)
            elif isinstance(it, Line):
                self.line(it)
            elif isinstance(it, Image):
                self.image(it)
            elif isinstance(it, TableGrid):
                self.table(it)


class PptxBackend:
    name = "pptx"

    def __init__(self, filename: str = "diagrams.pptx", split: bool = False, svg: str | None = None):
        self.filename = filename
        self.split = split
        self.svg = (svg or os.environ.get("SPINAGRAM_PPTX_SVG") or "vector").lower()
        if self.svg not in ("vector", "raster"):
            raise ValueError(f"PptxBackend(svg=...) must be 'vector' or 'raster', not {self.svg!r}")

    @staticmethod
    def _deck(w: float, h: float):
        from pptx import Presentation

        prs = Presentation()
        prs.slide_width, prs.slide_height = _e(w), _e(h)
        return prs

    def render(self, scenes: Sequence[Scene], out_dir: Path) -> list[Path]:
        _require()
        out_dir.mkdir(parents=True, exist_ok=True)
        if not scenes:
            return []
        if self.split:
            paths = []
            for sc in scenes:
                prs = self._deck(sc.width, sc.height)
                _Slide(prs, sc, self.svg).build()
                p = out_dir / f"{sc.slug}.pptx"
                prs.save(p)
                paths.append(p)
            return paths
        prs = self._deck(scenes[0].width, scenes[0].height)
        svg_parts: dict = {}
        for sc in scenes:
            _Slide(prs, sc, self.svg, svg_parts).build()
        p = out_dir / self.filename
        prs.save(p)
        return [p]
